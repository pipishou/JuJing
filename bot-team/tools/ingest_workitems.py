#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""工作任务清单 → workitems.v1 快照（工作任务管理 pack 的数据入口）。

输入：企微文档「工作事项清单」导出的 CSV / XLSX，或手工维护的同构表。
输出：workitems.v1 JSON（summary + items + provenance），供
      - 场景卡 current_value（逾期率/在办数/阻塞数）
      - 报告 data_refs（带 sha256，可溯源）
      - 决策卡「砍/延/升级」建议的排他性证据

为什么单独一个工具：清单是「行 = 事项」的快照，不是「行 = 日期」的时间序列，
generic ingest（tools/ingest_task_data.py）按 T/T30 折线设计，套不上事项级语义。

用法：
  python -B tools/ingest_workitems.py 清单.csv --owner 本人 -o out.json
  python -B tools/ingest_workitems.py 清单.xlsx --as-of 2026-10-28 --field-map fm.yaml
缺列/空表/日期非法一律 rc=2 并单行 JSON 报错（不静默补默认值）。
"""
import argparse
import csv
import hashlib
import json
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

SCHEMA = "workitems.v1"
REQUIRED_COLUMNS = ("title", "owner", "due", "status")

# 语义列 → 表头候选（顺序即优先级；用户表头千奇百怪，先认中文再认英文）
COLUMN_ALIASES = {
    "title": ("工作事项", "事项", "任务", "事项名称", "工作内容", "title", "item", "task"),
    "owner": ("负责人", "责任人", "执行人", "owner", "assignee"),
    "due": ("截止日", "截止日期", "计划完成", "完成时间", "due", "due_date", "deadline"),
    "status": ("状态", "进展", "当前状态", "status", "state"),
    "priority": ("优先级", "重要度", "priority"),
    "hours": ("预估工时", "工时", "预计工时", "hours", "estimate"),
    "blocker": ("阻塞", "阻塞原因", "卡点", "风险", "blocker", "risk"),
    "source": ("来源", "提出人", "source"),
    "done_date": ("完成日", "实际完成", "done_date", "closed_at"),
}

STATUS_DONE = ("已完成", "完成", "已关闭", "关闭", "done", "closed", "finished")
STATUS_BLOCKED = ("阻塞", "卡住", "blocked", "waiting")
STATUS_OPEN = ("进行中", "在办", "待办", "未开始", "open", "todo", "doing", "in_progress")

DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d", "%Y-%m-%d %H:%M:%S")


class IngestError(Exception):
    """清单不可用（缺列/空表/日期非法），一律 rc=2。"""


def _norm(s):
    return re.sub(r"\s+", "", str(s or "")).lower()


def parse_date(value, field="date"):
    """容忍 Excel 序列号 / 多种分隔符；识别不了就报错，不猜。"""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    # Excel 数字日期（45000 附近）；只认 20000-60000 区间，避免把「3」当日期
    if re.fullmatch(r"\d{5}(\.\d+)?", text):
        serial = float(text)
        if 20000 <= serial <= 60000:
            return (date(1899, 12, 30) + timedelta(days=int(serial)))
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise IngestError(f"{field} 日期无法识别: {text!r}")


def status_kind(raw):
    text = _norm(raw)
    if not text:
        return "unknown"
    for token in STATUS_DONE:
        if _norm(token) in text:
            return "done"
    for token in STATUS_BLOCKED:
        if _norm(token) in text:
            return "blocked"
    for token in STATUS_OPEN:
        if _norm(token) in text:
            return "open"
    return "unknown"


def read_rows(path, field_map=None):
    """读 CSV / XLSX → [(sheet, header, rows)]；field_map.source_columns 可覆盖表头。"""
    path = Path(path)
    if not path.is_file():
        raise IngestError(f"文件不存在: {path}")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        raw = path.read_bytes()
        text = None
        for enc in ("utf-8-sig", "gb18030", "utf-8"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise IngestError("CSV 编码无法识别（试过 utf-8-sig/gb18030/utf-8）")
        reader = csv.reader(text.splitlines())
        rows = [r for r in reader if any((c or "").strip() for c in r)]
        if not rows:
            raise IngestError("CSV 是空表")
        return [("csv", rows[0], rows[1:])]
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook  # noqa: PLC0415
        except ImportError as exc:  # pragma: no cover
            raise IngestError("读 xlsx 需要 openpyxl（venv 里已有；pip install openpyxl）") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        out = []
        try:
            for ws in wb.worksheets:
                rows = [list(r) for r in ws.iter_rows(values_only=True)]
                rows = [r for r in rows if any((c is not None and str(c).strip()) for c in r)]
                if rows:
                    out.append((ws.title, rows[0], rows[1:]))
        finally:
            wb.close()
        if not out:
            raise IngestError("工作簿里没有非空工作表")
        return out
    if suffix == ".xls":
        raise IngestError("旧的 .xls 请另存为 .xlsx 或 CSV 后再导入")
    raise IngestError(f"不支持的文件类型: {suffix}（支持 .csv/.xlsx/.xlsm）")


def map_columns(header, field_map=None):
    """表头 → 语义列；返回 {semantic: index}，缺必填列抛错。"""
    override = {}
    if isinstance(field_map, dict):
        src = field_map.get("source_columns") or field_map
        if isinstance(src, dict):
            override = {str(k): str(v) for k, v in src.items()}
    cells = [_norm(c) for c in header]
    mapping = {}
    for semantic, aliases in COLUMN_ALIASES.items():
        want = override.get(semantic)
        if want:
            target = _norm(want)
            if target in cells:
                mapping[semantic] = cells.index(target)
                continue
            raise IngestError(f"field_map 指定的列在表头里找不到: {semantic} -> {want!r}")
        for alias in aliases:
            key = _norm(alias)
            if key in cells:
                mapping[semantic] = cells.index(key)
                break
    missing = [c for c in REQUIRED_COLUMNS if c not in mapping]
    if missing:
        raise IngestError(
            "清单缺必填列: " + ", ".join(missing)
            + "；表头是 " + ", ".join(str(c) for c in header))
    return mapping


def _cell(row, mapping, semantic):
    idx = mapping.get(semantic)
    if idx is None or idx >= len(row):
        return None
    value = row[idx]
    if isinstance(value, str):
        value = value.strip()
    return value if value not in ("", None) else None


def build_snapshot(path, as_of=None, owner=None, field_map=None):
    """清单 → workitems.v1 dict（不写盘）。"""
    path = Path(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    as_of_date = parse_date(as_of, "as_of") if as_of else date.today()
    items = []
    sheets = read_rows(path, field_map)
    for sheet, header, rows in sheets:
        mapping = map_columns(header, field_map)
        for row_no, row in enumerate(rows, start=2):
            title = _cell(row, mapping, "title")
            if not title:
                continue
            item_owner = _cell(row, mapping, "owner")
            if owner and item_owner and _norm(item_owner) != _norm(owner):
                continue
            due = parse_date(_cell(row, mapping, "due"), f"{sheet}!{row_no} 截止日")
            done_date = parse_date(_cell(row, mapping, "done_date"), f"{sheet}!{row_no} 完成日")
            raw_status = _cell(row, mapping, "status")
            kind = status_kind(raw_status)
            hours = _cell(row, mapping, "hours")
            try:
                hours = float(str(hours).replace("h", "").strip()) if hours not in (None, "") else None
            except ValueError:
                hours = None
            items.append({
                "id": f"{sheet}-{row_no}",
                "title": str(title),
                "owner": str(item_owner) if item_owner else "",
                "due": due.isoformat() if due else None,
                "done_date": done_date.isoformat() if done_date else None,
                "status_raw": str(raw_status) if raw_status else "",
                "status": kind,
                "priority": str(_cell(row, mapping, "priority") or ""),
                "hours": hours,
                "blocker": str(_cell(row, mapping, "blocker") or ""),
                "source": str(_cell(row, mapping, "source") or ""),
                "overdue_days": (as_of_date - due).days if (due and kind != "done" and due < as_of_date) else 0,
            })
    if not items:
        raise IngestError("清单里没有可用事项行（title 全空或全部被 --owner 过滤掉）")
    summary = summarize(items, as_of_date)
    return {
        "schema": SCHEMA,
        "as_of": as_of_date.isoformat(),
        "owner_filter": owner or "",
        "source": {
            "path": str(path),
            "name": path.name,
            "sha256": digest,
            "size": path.stat().st_size,
            "sheets": [s[0] for s in sheets],
            "rows": len(items),
        },
        "summary": summary,
        "items": items,
    }


def summarize(items, as_of_date):
    """统计口径写死在 docstring，避免「数字对不上」的扯皮。

    - open_items：status ∈ {open, blocked, unknown}（未标记完成即视为在办）
    - overdue_items：未完成且 due < as_of
    - due_this_week：未完成且 due 在 [as_of, as_of+7d]
    - overdue_rate：overdue/在办（分母为 0 记 null，不写 0 冒充好成绩）
    - closure_rate：done 且 done_date ≤ due 的条数 / done 条数
    - data_gaps：缺负责人 / 缺截止日 的条数（评分卡完整性维度的证据）
    """
    week_end = as_of_date + timedelta(days=7)
    open_items = [i for i in items if i["status"] != "done"]
    done_items = [i for i in items if i["status"] == "done"]
    overdue = [i for i in open_items if i["overdue_days"] > 0]
    due_week = [i for i in open_items if i["due"] and as_of_date <= date.fromisoformat(i["due"]) <= week_end]
    blocked = [i for i in open_items if i["status"] == "blocked" or i["blocker"]]
    no_owner = [i for i in items if not i["owner"]]
    no_due = [i for i in items if not i["due"]]
    on_time = 0
    for item in done_items:
        if item["due"] and item["done_date"] and item["done_date"] <= item["due"]:
            on_time += 1
    return {
        "total_items": len(items),
        "open_items": len(open_items),
        "overdue_items": len(overdue),
        "due_this_week": len(due_week),
        "blocked_items": len(blocked),
        "done_items": len(done_items),
        "overdue_rate": round(len(overdue) * 100.0 / len(open_items), 2) if open_items else None,
        "closure_rate": round(on_time * 100.0 / len(done_items), 2) if done_items else None,
        "data_gaps": {"no_owner": len(no_owner), "no_due": len(no_due)},
        "worst_three": [
            {k: i[k] for k in ("title", "owner", "due", "overdue_days", "blocker")}
            for i in sorted(overdue, key=lambda x: -x["overdue_days"])[:3]
        ],
    }


def load_field_map(path):
    """读 field-map：优先 YAML（PyYAML），缺库时按 JSON 解析。"""
    raw = Path(path).read_text(encoding="utf-8-sig")
    try:
        import yaml  # noqa: PLC0415
        doc = yaml.safe_load(raw)
    except Exception:  # noqa: BLE001
        doc = json.loads(raw)
    if not isinstance(doc, dict):
        raise IngestError("field-map 顶层必须是对象")
    return doc


def main(argv=None):
    ap = argparse.ArgumentParser(description="工作任务清单 → workitems.v1 快照")
    ap.add_argument("path", help="清单文件（.csv / .xlsx）")
    ap.add_argument("-o", "--out", help="输出 JSON 路径（缺省 stdout）")
    ap.add_argument("--as-of", help="统计基准日（默认今天）")
    ap.add_argument("--owner", help="只统计该负责人的事项")
    ap.add_argument("--field-map", help="列名映射 YAML/JSON：{source_columns: {title: 事项名称, ...}}")
    args = ap.parse_args(argv)
    try:
        field_map = load_field_map(args.field_map) if args.field_map else None
        doc = build_snapshot(args.path, as_of=args.as_of, owner=args.owner, field_map=field_map)
    except IngestError as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False))
        return 2
    payload = json.dumps(doc, ensure_ascii=False, indent=2)
    if args.out:
        Path(args.out).write_text(payload + "\n", encoding="utf-8")
        print(json.dumps({"ok": True, "out": args.out, "summary": doc["summary"]}, ensure_ascii=False))
    else:
        print(payload)
    return 0


if __name__ == "__main__":
    sys.exit(main())
